import copy
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
from evals.core.dataset import export_dataset, import_dataset, freeze_row, load_dataset
from evals.core.runner import create_run, run_cases
from evals.core.evaluator import score_run
from evals.core.report import report_run, compare_runs
from evals.core.storage import read_json, read_jsonl, atomic_json
from evals.scenarios import SCENARIOS, Scenario


def fixture():
    messages = [{"messageId": "root", "createdAt": "2026-09-24T00:00:00Z", "senderRole": "user", "text": "订单失败 =SUM(1,2)", "hasArtifact": False},
                {"messageId": "within", "createdAt": "2026-09-24T00:04:59Z", "senderRole": "user", "text": "请排查", "replyTo": "root", "hasArtifact": False},
                {"messageId": "future", "createdAt": "2026-09-24T00:05:00Z", "senderRole": "user", "text": "FUTURE_SECRET", "hasArtifact": False}]
    return {"ticket": {"baseId": "base", "tableId": "table", "recordId": "rec-1"}, "number": "1", "title": "=SECRET_TITLE",
            "history_complete": True, "snapshot_version": 1, "messages": messages,
            "timing": [{"messageId": m["messageId"], "createdAt": m["createdAt"], **({"rootId": "root", "parentId": "root"} if i else {})} for i, m in enumerate(messages)],
            "shadow": {"analysis": {"analysis": {"intent": {"summary": "SHADOW_SECRET"}}}}}


class PipelineTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.addCleanup(self.temp.cleanup); self.root = Path(self.temp.name)

    def annotated(self, rows=None, both=True):
        export = export_dataset(self.root, rows or [fixture()], "synthetic")
        excel = self.root / "exports" / export / "annotations.xlsx"
        book = load_workbook(excel); sheet = book.active; columns = {c.value: c.column for c in sheet[1]}
        for row in range(2, sheet.max_row + 1):
            values = {"分析_纳入": "是", "分析_审核": "已审核", "分析_意图": "troubleshoot", "分析_摘要": "GOLD_SECRET",
                      "审核人": "reviewer", "答复_纳入": "是" if both else "否", "答复_审核": "已审核", "答复_参考答复": "请提供报错"}
            for name, value in values.items(): sheet.cell(row, columns[name], value)
        book.save(excel); book.close()
        return excel

    def dataset(self):
        return import_dataset(self.root, self.annotated())["dataset_id"]

    def fake_bridge(self, operation, data, **kwargs):
        if operation == "defaults": return {"prompts": {"lark_ticket.shadow.summarize": "{{ticket_context}}"}}
        serialized = str(data)
        for secret in ("GOLD_SECRET", "SHADOW_SECRET", "FUTURE_SECRET", "SECRET_TITLE"):
            self.assertNotIn(secret, serialized)
        self.assertNotIn("expected", data)
        return {"status": "ok", "output": {"analysis": {"intent": {"intentType": "troubleshoot"}}, "summary": "需要排查"}}

    def measure(self, *args):
        return [{"name": "摘要准确完整", "score": .8, "passed": True, "reason": "mock", "status": "ok"}]

    def make_run(self, dataset):
        run = create_run(self.root, dataset, "ticket.early_analysis", {"provider": "deepseek", "model": "test"}, self.fake_bridge)
        run_cases(self.root, run, self.fake_bridge)
        return run

    def test_window_boundary_and_invalid_histories(self):
        frozen = freeze_row(fixture())
        self.assertEqual([m["messageId"] for m in frozen["input"]["messages"]], ["root", "within"])
        for mutation in ("history", "missing_root", "edited", "timestamp", "deleted"):
            row = fixture()
            if mutation == "history": row["history_complete"] = False
            if mutation == "missing_root": row["timing"].pop(0)
            if mutation == "edited": row["timing"][0]["updatedAt"] = "2026-09-24T00:10:00Z"
            if mutation == "timestamp": row["timing"][0]["createdAt"] = "invalid"
            if mutation == "deleted": row["timing"][0]["deleted"] = True
            self.assertFalse(freeze_row(row)["eligible"], mutation)

    def test_import_two_cases_idempotent_and_independent_from_export(self):
        excel = self.annotated()
        first = import_dataset(self.root, excel); second = import_dataset(self.root, excel)
        self.assertEqual(first["dataset_id"], second["dataset_id"]); self.assertEqual(first["included"], 2)
        import shutil
        shutil.rmtree(excel.parent)
        _, cases, inputs = load_dataset(self.root, first["dataset_id"])
        self.assertEqual(len(cases), 2); self.assertEqual(len(inputs), 1)

    def test_independent_review_blocks_entire_import(self):
        excel = self.annotated(); book = load_workbook(excel); sheet = book.active
        col = {c.value: c.column for c in sheet[1]}; sheet.cell(2, col["答复_审核"], "待审核"); book.save(excel); book.close()
        with self.assertRaisesRegex(ValueError, "review"): import_dataset(self.root, excel)
        self.assertFalse((self.root / "datasets").exists())

    def test_changed_display_input_rejected_and_formulas_literal(self):
        excel = self.annotated(); book = load_workbook(excel); sheet = book.active
        col = {c.value: c.column for c in sheet[1]}
        self.assertEqual(sheet.cell(2, col["标题"]).data_type, "s")
        sheet.cell(2, col["窗口消息_1"], "tampered"); book.save(excel); book.close()
        with self.assertRaisesRegex(ValueError, "window text"): import_dataset(self.root, excel)

    def test_long_messages_and_emoji_roundtrip(self):
        row = fixture(); row["messages"][0]["text"] = "🙂中文\n" * 10000
        excel = self.annotated([row]); self.assertEqual(import_dataset(self.root, excel)["included"], 2)

    def test_empty_and_duplicate_and_hash_rejection(self):
        export = export_dataset(self.root, [fixture()], "synthetic")
        path = self.root / "exports" / export
        with self.assertRaisesRegex(ValueError, "No reviewed"): import_dataset(self.root, path / "annotations.xlsx")
        (path / "inputs.jsonl").write_text("{}\n")
        with self.assertRaisesRegex(ValueError, "hash"): import_dataset(self.root, path / "annotations.xlsx")
        with self.assertRaisesRegex(ValueError, "Duplicate"): export_dataset(self.root, [fixture(), fixture()], "test")

    def test_full_pipeline_resume_report_comparison_and_no_leak(self):
        dataset = self.dataset(); run = self.make_run(dataset)
        def forbidden(*args, **kwargs): raise AssertionError("resume regenerated result")
        run_cases(self.root, run, forbidden)
        score = score_run(self.root, run, {"model": "mock"}, self.measure)
        report = report_run(self.root, run, score); self.assertTrue(Path(report).exists())
        summary = read_json(Path(report).parent / "summary.json")
        self.assertEqual(summary["planned"], 1); self.assertEqual(summary["passed"], 1)
        run2 = self.make_run(dataset); score2 = score_run(self.root, run2, {"model": "mock"}, self.measure)
        compare = compare_runs(self.root, run, run2, score, score2)
        self.assertEqual(read_json(self.root / "comparisons" / compare / "manifest.json")["comparable"], 1)
        with self.assertRaisesRegex(ValueError, "Specify"):
            score_run(self.root, run, {"model": "mock"}, self.measure); report_run(self.root, run)

    def test_failure_counts_and_rule_mismatch(self):
        dataset = self.dataset(); run = self.make_run(dataset)
        def fail(*args): raise RuntimeError("provider secret")
        score = score_run(self.root, run, {"model": "mock"}, fail)
        report = report_run(self.root, run, score)
        self.assertEqual(read_json(Path(report).parent / "summary.json")["judge_error"], 1)
        score2 = score_run(self.root, run, {"model": "other"}, self.measure)
        compare = compare_runs(self.root, run, run, score, score2)
        self.assertEqual(read_json(self.root / "comparisons" / compare / "manifest.json")["comparable"], 0)

    def test_changed_run_output_cannot_be_reported_with_old_score(self):
        run = self.make_run(self.dataset()); score = score_run(self.root, run, {"model": "mock"}, self.measure)
        path = next((self.root / "runs" / run / "cases").glob("*.json")); actual = read_json(path)
        actual["output"] = "changed"; atomic_json(path, actual)
        with self.assertRaisesRegex(ValueError, "changed since scoring"): report_run(self.root, run, score)

    def test_config_rejects_secrets(self):
        with self.assertRaisesRegex(ValueError, "Unknown model config"):
            create_run(self.root, self.dataset(), "ticket.early_analysis", {"provider": "deepseek", "model": "test", "apiKey": "secret"}, self.fake_bridge)

    def test_third_scenario_reuses_export_import_run_score_report(self):
        third = Scenario("ticket.synthetic", "第三", 1, ("answer",), (("合成指标", "按参考内容判断"),),
                         expected_columns=(("answer", "答案"),), output_field="answerMarkdown")
        with patch.dict(SCENARIOS, {third.id: third}):
            excel = self.annotated(both=False)
            book = load_workbook(excel); sheet = book.active; columns = {c.value: c.column for c in sheet[1]}
            for key, value in {"分析_纳入": "否", "第三_纳入": "是", "第三_审核": "已审核", "第三_答案": "合成"}.items():
                sheet.cell(2, columns[key], value)
            book.save(excel); book.close()
            dataset = import_dataset(self.root, excel)["dataset_id"]
            run = create_run(self.root, dataset, third.id, {"provider": "deepseek", "model": "test"}, self.fake_bridge)
            run_cases(self.root, run, lambda *args, **kwargs: {"status": "ok", "output": {"answerMarkdown": "合成"}})
            score = score_run(self.root, run, {"model": "mock"}, self.measure)
            self.assertTrue(Path(report_run(self.root, run, score)).exists())

    def test_resume_after_interruption_and_generation_failures(self):
        second = fixture(); second["ticket"]["recordId"] = "rec-2"
        dataset = import_dataset(self.root, self.annotated([fixture(), second], both=False))["dataset_id"]
        run = create_run(self.root, dataset, "ticket.early_analysis", {"provider": "deepseek", "model": "test"}, self.fake_bridge)
        count = 0
        def interrupted(*args, **kwargs):
            nonlocal count
            count += 1
            if count == 2: raise KeyboardInterrupt()
            return {"status": "invalid_output", "error": {"errorCode": "SHADOW_OUTPUT_INVALID"}}
        with self.assertRaises(KeyboardInterrupt): run_cases(self.root, run, interrupted)
        self.assertEqual(len(list((self.root / "runs" / run / "cases").glob("*.json"))), 1)
        run_cases(self.root, run, lambda *args, **kwargs: {"status": "error", "error": {"errorCode": "CONFIG_MISSING"}})
        score = score_run(self.root, run, {"model": "mock"}, self.measure)
        summary = read_json(Path(report_run(self.root, run, score)).parent / "summary.json")
        self.assertEqual((summary["planned"], summary["quality_failed"], summary["generation_error"], summary["scoring_ok"]), (2, 1, 1, 0))

    def test_changed_annotations_are_new_version_and_not_comparable(self):
        excel = self.annotated(); dataset = import_dataset(self.root, excel)["dataset_id"]
        baseline = self.make_run(dataset); a = score_run(self.root, baseline, {"model": "mock"}, self.measure)
        book = load_workbook(excel); sheet = book.active; cols = {c.value: c.column for c in sheet[1]}
        sheet.cell(2, cols["分析_摘要"], "changed golden"); book.save(excel); book.close()
        updated = import_dataset(self.root, excel)["dataset_id"]; self.assertNotEqual(dataset, updated)
        candidate = self.make_run(updated); b = score_run(self.root, candidate, {"model": "mock"}, self.measure)
        comparison = compare_runs(self.root, baseline, candidate, a, b)
        self.assertEqual(read_json(self.root / "comparisons" / comparison / "manifest.json")["comparable"], 0)

    def test_snapshot_of_knowledge_is_independent_and_verified(self):
        from evals.core.runner import freeze_wiki, verify_wiki
        source = self.root / "source/docs/llm-wiki"; source.mkdir(parents=True)
        (source / "page.md").write_text("original")
        run = self.root / "run"
        files = freeze_wiki(self.root / "source", run / "knowledge")
        (source / "page.md").write_text("changed")
        verify_wiki(run, {"knowledge_files": files})
        (run / "knowledge/docs/llm-wiki/page.md").write_text("tampered")
        with self.assertRaisesRegex(ValueError, "knowledge modified"): verify_wiki(run, {"knowledge_files": files})


if __name__ == "__main__": unittest.main()
